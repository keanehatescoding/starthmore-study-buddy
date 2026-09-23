"""Phase 2 tests: extraction dispatch + chunking (FakeLLM, no network)."""

import io

import pytest
from sqlmodel import Session, SQLModel, create_engine, select

from app.chunk import chunk_resource, locate, presplit
from app.extract import (
    ExtractError,
    SkipResource,
    extract_bytes,
    extract_docx,
    extract_pptx,
    extract_resource_text,
)
from app.models import Chunk, Course, Resource, Topic
from app.pipeline import run_chunking, run_extraction


class FakeLLM:
    def complete_json(self, system, user):
        # verbatim halves of the section -> offsets must resolve
        text = user.split("\n\n", 1)[1]
        half = len(text) // 2
        return {"chunks": [
            {"title": "C0", "content": text[:half]},
            {"title": "C1", "content": text[half:]},
        ]}


@pytest.fixture()
def session():
    engine = create_engine("sqlite:///:memory:")
    SQLModel.metadata.create_all(engine)
    with Session(engine) as s:
        yield s


def _resource(session, **kw):
    course = session.exec(
        select(Course).where(Course.source_id == "c1")).first()
    if course is None:
        course = Course(source="moodle", source_id="c1", name="C")
        session.add(course)
        session.commit()
    topic = session.exec(
        select(Topic).where(Topic.course_id == course.id)).first()
    if topic is None:
        topic = Topic(course_id=course.id, source_id="t1", title="T")
        session.add(topic)
        session.commit()
    kw.setdefault("topic_id", topic.id)
    kw.setdefault("source", "moodle")
    kw.setdefault("source_id", "r1")
    kw.setdefault("type", "file")
    kw.setdefault("title", "R")
    kw.setdefault("status", "pending")
    r = Resource(**kw)
    session.add(r)
    session.commit()
    session.refresh(r)
    return r


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello docx world")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 10, 10)
    box.text_frame.text = "Hello pptx world"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_docx_roundtrip():
    assert "Hello docx world" in extract_docx(_docx_bytes())


def test_pptx_roundtrip():
    text = extract_pptx(_pptx_bytes())
    assert "Slide 1" in text and "Hello pptx world" in text


def test_dispatch_and_unsupported():
    assert "Hello" in extract_bytes(_docx_bytes(), None, "notes.docx")
    assert "plain" in extract_bytes(b"plain text", "text/plain", "n.txt")
    with pytest.raises(ExtractError):
        extract_bytes(b"\x00\x01", "application/x-unknown", "n.bin")


def test_video_without_url_skipped(session):
    r = _resource(session, type="video", raw_url=None)
    with pytest.raises(SkipResource):
        extract_resource_text(r)


def test_link_skipped(session):
    r = _resource(session, type="link", raw_url="https://example.com")
    with pytest.raises(SkipResource):
        extract_resource_text(r)


def test_presplit_and_locate():
    text = ("para one\n\n" * 5000)
    assert len(presplit(text)) > 1
    assert locate("b", "abc") == (1, 2)
    assert locate("zzz", "abc") == (None, None)


def test_locate_fuzzy_whitespace():
    full = "Identifying  concepts\nrelated  to  networks"
    assert locate("Identifying concepts related to networks", full) == (0, 43)


def test_chunk_short_text_no_llm(session):
    r = _resource(session, status="extracted", extracted_text="tiny but real")
    assert chunk_resource(session, r, None) == 1
    chunk = select(Chunk)
    assert len(session.exec(chunk).all()) == 1


def test_chunk_regenerates_on_change(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500)
    assert chunk_resource(session, r, FakeLLM()) == 2
    r.status = "pending"  # sync saw new content
    session.add(r)
    session.commit()
    assert chunk_resource(session, r, FakeLLM()) == 2
    assert len(session.exec(select(Chunk)).all()) == 2  # replaced, not doubled


def test_chunk_cached(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500)
    assert chunk_resource(session, r, FakeLLM()) == 2
    assert chunk_resource(session, r, FakeLLM()) == 0


def test_run_extraction_counts(session):
    _resource(session, source_id="ok", type="page_text", text="page content here",
              extracted_text="page content here")
    _resource(session, source_id="skip", type="link", raw_url="https://example.com")
    counts = run_extraction(session, downloader=None)
    assert counts == {"extracted": 1, "skipped": 1, "failed": 0}


def test_run_chunking_counts(session):
    _resource(session, source_id="a", status="extracted", extracted_text="short one")
    counts = run_chunking(session, None)
    assert counts["chunks"] == 1 and counts["resources"] == 1
